import io
import pandas as pd
import numpy as np
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generar_excel_plan(df_res, horas_por_linea=None):
    """
    Genera un archivo Excel (.xlsx) estructurado en bloques:
    - Pestaña Resumen_Cobertura_SKUs
    - Pestaña Única: Secuenciamiento_Global (Bloques divididos visualmente por línea)
    - Pestaña consolidada Plan_General_Sugerido
    - Pestaña de alertas SKUs_Exceden_Capacidad
    """
    output = io.BytesIO()
    df_export = df_res.copy()
    
    # 1. MÉTRICAS DE COBERTURA BASE 30 DÍAS
    demanda_diaria = np.where(df_export['Demanda_Mensual_Clean'] > 0, df_export['Demanda_Mensual_Clean'] / 30.0, 0.0001)
    df_export['Días de Inventario Actual'] = df_export['Inventario_Actual'] / demanda_diaria
    df_export['Días de Inventario Final'] = (df_export['Inventario_Actual'] + df_export['Sugerido_Final_Ajustado']) / demanda_diaria
    
    df_export.loc[df_export['Demanda_Mensual_Clean'] <= 0, 'Días de Inventario Actual'] = 999.0
    df_export.loc[df_export['Demanda_Mensual_Clean'] <= 0, 'Días de Inventario Final'] = 999.0

    df_prog = df_export[df_export['Sugerido_Final_Ajustado'] > 0].copy()
    dias_nombres = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado']
    dias_programados = []
    
    # 2. CÁLCULO DE DÍAS ASIGNADOS
    for idx, row in df_prog.iterrows():
        h_ini = row['Hora_Inicio_Proyectada']
        h_fin = row['Hora_Fin_Proyectada']
        linea = row['Linea_Produccion']
        
        h_bruta_linea = horas_por_linea.get(str(linea), 120.0) if horas_por_linea else 120.0
        h_por_dia = max(1.0, h_bruta_linea / 6.0)
        
        idx_dia_ini = min(5, int(h_ini // h_por_dia))
        idx_dia_fin = min(5, int(max(0.0, h_fin - 0.1) // h_por_dia))
        
        if idx_dia_ini == idx_dia_fin:
            dias_programados.append(dias_nombres[idx_dia_ini])
        else:
            dias_programados.append(f"{dias_nombres[idx_dia_ini]} a {dias_nombres[idx_dia_fin]}")
            
    df_prog['Día_Programado'] = dias_programados
    
    cols_diarias = [
        'Orden_Secuencia', 'Día_Programado', 'Linea_Produccion', 'SKU', 'DESCRIPCION', 'FORMATO',
        'Prioridad_Nivel', 'Hora_Inicio_Proyectada', 'Hora_Fin_Proyectada', 'Sugerido_Final_Ajustado',
        'Horas_Asignadas', 'Horas_Setup_Formato', 'Horas_Setup_Referencia', 'Estado_Produccion'
    ]
    
    cols_resumen = [
        'Linea_Produccion', 'SKU', 'DESCRIPCION', 'Inventario_Actual', 
        'Días de Inventario Actual', 'Sugerido_Final_Ajustado', 
        'Días de Inventario Final', 'Estado_Produccion'
    ]
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        
        # Pestañas tradicionales gestionadas vía Pandas (Resumen, General y Excedidos)
        df_resumen = df_export[cols_resumen].sort_values(by=['Linea_Produccion', 'Días de Inventario Actual'])
        df_resumen.to_excel(writer, index=False, sheet_name='Resumen_Cobertura_SKUs')
        df_export.to_excel(writer, index=False, sheet_name='Plan_General_Sugerido')
        
        df_excedidos = df_export[df_export['Estado_Produccion'] == 'Excede Capacidad']
        if len(df_excedidos) > 0:
            df_excedidos.to_excel(writer, index=False, sheet_name='SKUs_Exceden_Capacidad')

        workbook = writer.book
        
        # 3. PESTAÑA PERSONALIZADA: SECUENCIAMIENTO_GLOBAL (Bloques Visuales)
        ws_sec = workbook.create_sheet('Secuenciamiento_Global', 1) # Insertar en la segunda posición
        
        # Definición de Estilos
        title_fill = PatternFill(start_color="10253F", end_color="10253F", fill_type="solid") # Azul Marino Oscuro
        title_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Azul Clásico
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        sugerido_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Verde Claro
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
        )
        
        current_row = 1
        lineas_unicas_prog = sorted(list(df_prog['Linea_Produccion'].astype(str).unique()))
        
        for linea_name in lineas_unicas_prog:
            df_linea_sec = df_prog[df_prog['Linea_Produccion'].astype(str) == str(linea_name)][cols_diarias]
            if df_linea_sec.empty:
                continue
                
            # A) Título del bloque (Divisor visual de la máquina)
            cell_title = ws_sec.cell(row=current_row, column=1)
            cell_title.value = f"SECUENCIAMIENTO: {str(linea_name).upper()}"
            cell_title.fill = title_fill
            cell_title.font = title_font
            ws_sec.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(cols_diarias))
            cell_title.alignment = Alignment(horizontal="center", vertical="center")
            current_row += 1
            
            # B) Cabeceras de columnas para este bloque
            col_sugerido_idx = None
            for col_idx, col_name in enumerate(cols_diarias, 1):
                cell_h = ws_sec.cell(row=current_row, column=col_idx)
                cell_h.value = col_name
                cell_h.fill = header_fill
                cell_h.font = header_font
                cell_h.border = thin_border
                cell_h.alignment = Alignment(horizontal="center", vertical="center")
                
                # Detectar qué índice tiene la columna sugerida
                if col_name == 'Sugerido_Final_Ajustado':
                    col_sugerido_idx = col_idx
            current_row += 1
            
            # C) Relleno de datos de la línea
            for _, row_data in df_linea_sec.iterrows():
                for col_idx, col_name in enumerate(cols_diarias, 1):
                    cell_d = ws_sec.cell(row=current_row, column=col_idx)
                    val = row_data[col_name]
                    
                    # Formateo visual
                    if isinstance(val, float):
                        cell_d.number_format = '#,##0.0' if col_name.startswith('Horas') else '#,##0'
                    elif isinstance(val, int):
                        cell_d.number_format = '#,##0'
                        
                    cell_d.value = val
                    cell_d.border = thin_border
                    
                    # Resaltar en verde claro
                    if col_idx == col_sugerido_idx:
                        cell_d.fill = sugerido_fill 
                current_row += 1
                
            # D) Dos filas en blanco antes de la próxima máquina
            current_row += 2
            
        # Ajuste de ancho de columnas para Secuenciamiento_Global
        for col in ws_sec.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.value:
                    val_len = len(str(cell.value))
                    if val_len > max_len: max_len = val_len
            ws_sec.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)
            
        # 4. APLICACIÓN DE ESTILOS AL RESTO DE PESTAÑAS (PANDAS)
        for sheet_name in workbook.sheetnames:
            if sheet_name == 'Secuenciamiento_Global':
                continue # Ya procesada manualmete
                
            worksheet = workbook[sheet_name]
            
            # Ubicar la columna Sugerido
            col_sugerido_idx = None
            for col_num in range(1, worksheet.max_column + 1):
                cell = worksheet.cell(row=1, column=col_num)
                if cell.value == 'Sugerido_Final_Ajustado':
                    col_sugerido_idx = col_num
                
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for col in worksheet.columns:
                col_letter = get_column_letter(col[0].column)
                max_len = 0
                
                for cell in col:
                    if cell.row > 1 and cell.column == col_sugerido_idx:
                        cell.fill = sugerido_fill # Verde Claro
                    
                    cell.border = thin_border
                    try:
                        val_len = len(str(cell.value or ''))
                        if val_len > max_len: max_len = val_len
                    except: pass
                
                worksheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    return output.getvalue()